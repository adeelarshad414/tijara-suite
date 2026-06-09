from __future__ import annotations

import csv
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_BREAK
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
SCREENSHOTS = DOCS / "screenshots"
MARKDOWN_OUT = DOCS / "SCREENSHOT_USER_GUIDE.md"
DOCX_OUT = DOCS / "Tijara_Suite_Screenshot_User_Guide.docx"

FIGURES = [
    {
        "title": "Login And Authenticated App Shell",
        "path": SCREENSHOTS / "platform-superadmin" / "odoo-login.png",
        "caption": "Start from the Odoo login page, select the `tijara_dev` database when prompted, and sign in with the role account assigned to your workflow.",
        "steps": [
            "Open `http://localhost:8069/web/login?db=tijara_dev`.",
            "Use the role account from `docs/TEST_CREDENTIALS.csv`.",
            "After login, confirm the top Odoo navigation and company switcher are visible.",
        ],
    },
    {
        "title": "Back Office App Shell",
        "path": SCREENSHOTS / "platform-superadmin" / "odoo-app-shell.png",
        "caption": "The authenticated Odoo shell is the entry point for Apps, POS, Inventory, Sales, Purchases, Accounting, SaaS control, analytics, and configuration menus.",
        "steps": [
            "Use the app switcher to open the required module.",
            "Use role-specific menus for daily work and manager approvals.",
            "Keep demo operations in the local `tijara_dev` database.",
        ],
    },
    {
        "title": "Kiosk Self Ordering",
        "path": SCREENSHOTS / "public-display" / "kiosk-display.png",
        "caption": "The kiosk route supports dine-in, takeaway, pickup, B2C/B2B pricing, PKR prices, customer details, payment selection, and checkout.",
        "steps": [
            "Open `/tijara/display/tijara-demo-kiosk?db=tijara_dev` on a kiosk or tablet.",
            "Choose the service mode and customer type.",
            "Tap products to add them to the cart, enter customer details when needed, and press Checkout.",
        ],
    },
    {
        "title": "Customer Display",
        "path": SCREENSHOTS / "public-display" / "customer-display.png",
        "caption": "The customer-facing display shows the live order reference, item lines, quantities, PKR subtotal, discount, tax, and total.",
        "steps": [
            "Open `/tijara/display/tijara-demo-customer-display?db=tijara_dev` on the customer-side display.",
            "Keep it full screen near the POS counter.",
            "Use it to verify cart lines, discounts, GST, and totals with the customer before payment.",
        ],
    },
    {
        "title": "Queue Display",
        "path": SCREENSHOTS / "public-display" / "queue-display.png",
        "caption": "The queue display shows waiting, preparing, ready, and called tickets in large, readable cards for restaurant, bakery, and pickup counters.",
        "steps": [
            "Open `/tijara/display/tijara-demo-queue-display?db=tijara_dev` on a public screen.",
            "Use queue tickets from kiosk or POS orders.",
            "Advance ticket state from the operator workflow so the display reflects pickup readiness.",
        ],
    },
    {
        "title": "Menu Board",
        "path": SCREENSHOTS / "public-display" / "menu-board.png",
        "caption": "The menu board displays active menu items and promotions for restaurant, bakery, and food-service verticals.",
        "steps": [
            "Open `/tijara/display/tijara-demo-menu-board?db=tijara_dev`.",
            "Maintain menu items, Urdu/English names, and prices in display content or products.",
            "Use active start/end dates for timed menu or campaign changes.",
        ],
    },
    {
        "title": "Deals And Promotions Board",
        "path": SCREENSHOTS / "public-display" / "deals-board.png",
        "caption": "The deals board highlights discount campaigns and promotion messages for in-store screens.",
        "steps": [
            "Open `/tijara/display/tijara-demo-deals-board?db=tijara_dev`.",
            "Create promotions with display visibility enabled.",
            "Review price, discount, and availability before running the campaign in store.",
        ],
    },
]

ROLE_GUIDANCE = [
    ("platform_superadmin", "Provision tenants, control SaaS plans, review protected release evidence, and manage platform configuration."),
    ("tenant_admin", "Configure company settings, feature flags, staff access, POS hardware, receipt templates, dashboards, and approvals."),
    ("cashier", "Run B2C/B2B checkout, bill discounts, barcode refund scans, receipt print, and offline queue retry."),
    ("inventory_manager", "Maintain products, low-stock/expiry alerts, racks, shelves, bins, warehouses, and bulk import/export."),
    ("accountant", "Review settlements, refunds, chargebacks, draft accounting moves, and FBR queue evidence."),
    ("restaurant_operator", "Operate dine-in, takeaway, pickup, kiosk orders, queue tickets, kitchen status, menu boards, and pickup screens."),
    ("public_display", "Run public display routes for kiosk, customer display, queue, menu, deals, and promotions."),
]

WORKFLOW_SECTIONS = [
    (
        "POS Checkout, Discounts, Refunds, And Print",
        [
            "Open POS from the app shell and select the active Tijara Demo POS session.",
            "Choose B2C or B2B when the feature is enabled for the tenant.",
            "Add products by search, barcode scanner, QR scanner, or touch product cards.",
            "Apply an overall bill discount by percentage or amount; the paired value recalculates automatically.",
            "Take payment, validate the sale, print the receipt, and use the receipt barcode for future refund/exchange lookup.",
        ],
    ),
    (
        "Inventory, Expiry, Shelves, And Bulk Data",
        [
            "Use product forms for B2C/B2B prices, Urdu names, GST category, barcode aliases, and quick-sale flags.",
            "Use inventory intelligence for low-stock alerts, expiry alerts, warehouses, store rooms, racks, shelves, and bins.",
            "Use CSV import/export for products, prices, stock, contacts, hardware devices, receipt templates, storage positions, and promotions.",
            "Review import validation messages before applying data to a tenant database.",
        ],
    ),
    (
        "Restaurant, Bakery, And Pickup Operations",
        [
            "Use dine-in, takeaway, and pickup modes on kiosk or POS flows.",
            "Send kiosk orders into queue tickets and linked POS payments when the profile has a POS register and payment method.",
            "Use queue display for public ticket status and kitchen/operator views for preparation stages.",
            "Use menu and deals boards for active promotions, food menus, bakery offers, and pickup announcements.",
        ],
    ),
    (
        "SaaS Controls And Tenant Operations",
        [
            "Enable or disable B2B sales, queue display, promotion display, customer display, kiosk, and vertical packs by tenant plan.",
            "Use tenant provisioning manifests for DNS, ingress, admin bootstrap, backup, monitoring, and smoke evidence.",
            "Keep secrets in `secrets/.env.secrets`; keep non-secret runtime configuration in `.env` and deploy templates.",
        ],
    ),
    (
        "Payments, FBR, Analytics, And Production Evidence",
        [
            "Use payment provider records for JazzCash, Easypaisa, Stripe, generic/manual webhooks, settlement imports, refunds, and chargebacks.",
            "Use FBR queue records for dry-run/live adapter evidence until certified-provider credentials are available.",
            "Use analytics dashboards for POS, inventory, queue, promotions, trends, history, KPI snapshots, charts, and report catalogs.",
            "Use protected release evidence, security scans, monitoring evidence, backups, restore drills, and load tests before production sign-off.",
        ],
    ),
]


def load_credentials():
    path = DOCS / "TEST_CREDENTIALS.csv"
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def md_image(path: Path) -> str:
    return path.relative_to(DOCS).as_posix()


def write_markdown():
    captured_at = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    lines = [
        "# Tijara Suite Screenshot User Guide",
        "",
        f"Generated from live local Odoo screenshots on {captured_at}.",
        "",
        "## Live Verification Snapshot",
        "",
        "- Database: `tijara_dev`",
        "- Localization: Pakistan country, PKR currency, GST 18% sales tax verified",
        "- Demo users: seeded from `docs/TEST_CREDENTIALS.csv`",
        "- Screenshot source: `docs/screenshots/INDEX.md`",
        "",
        "## Demo Roles",
        "",
        "| Role | Login | Primary use |",
        "|---|---|---|",
    ]
    credentials = {row["persona"]: row["email"] for row in load_credentials()}
    for role, guidance in ROLE_GUIDANCE:
        lines.append(f"| `{role}` | `{credentials.get(role, '')}` | {guidance} |")

    lines.extend(["", "## Screenshot Walkthrough", ""])
    for figure in FIGURES:
        lines.extend(
            [
                f"### {figure['title']}",
                "",
                figure["caption"],
                "",
                f"![{figure['title']}]({md_image(figure['path'])})",
                "",
            ]
        )
        for index, step in enumerate(figure["steps"], 1):
            lines.append(f"{index}. {step}")
        lines.append("")

    lines.extend(["## Workflow Checklist", ""])
    for title, steps in WORKFLOW_SECTIONS:
        lines.extend([f"### {title}", ""])
        for step in steps:
            lines.append(f"- {step}")
        lines.append("")

    lines.extend(
        [
            "## Production Readiness Notes",
            "",
            "- Physical printer, scanner, drawer, scale, and display certification still needs real device evidence.",
            "- FBR live operation still needs certified-provider credentials and compliance sign-off.",
            "- Payment providers still need PSP certification, settlement reconciliation, refunds, and chargeback sign-off.",
            "- Full staging browser E2E should run against seeded users and real staging URLs before customer deployment.",
            "- Monitoring, alerting, restore drills, load tests, and security scans should be attached to release sign-off.",
            "",
        ]
    )
    MARKDOWN_OUT.write_text("\n".join(lines), encoding="utf-8")


def configure_doc(document: Document):
    section = document.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = document.styles
    styles["Normal"].font.name = "Calibri"
    styles["Normal"].font.size = Pt(11)
    for style_name, size, color in [
        ("Heading 1", 16, RGBColor(0x2E, 0x74, 0xB5)),
        ("Heading 2", 13, RGBColor(0x2E, 0x74, 0xB5)),
        ("Heading 3", 12, RGBColor(0x1F, 0x4D, 0x78)),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = color


def add_bullet(document: Document, text: str):
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.add_run(text)


def add_number(document: Document, text: str):
    paragraph = document.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.add_run(text)


def add_caption(document: Document, text: str):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(9)
    run = paragraph.add_run(text)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)


def add_picture(document: Document, path: Path):
    if not path.exists():
        add_bullet(document, f"Screenshot missing: {path.relative_to(ROOT)}")
        return
    document.add_picture(str(path), width=Inches(6.6))


def add_credentials_table(document: Document):
    credentials = {row["persona"]: row["email"] for row in load_credentials()}
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    header = table.rows[0].cells
    header[0].text = "Role"
    header[1].text = "Login"
    header[2].text = "Primary use"
    for role, guidance in ROLE_GUIDANCE:
        cells = table.add_row().cells
        cells[0].text = role
        cells[1].text = credentials.get(role, "")
        cells[2].text = guidance


def write_docx():
    document = Document()
    configure_doc(document)

    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(6)
    run = title.add_run("Tijara Suite Screenshot User Guide")
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(0x12, 0x26, 0x20)

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    subtitle.add_run("Live local Odoo walkthrough for POS, kiosk, displays, inventory, back office, SaaS controls, analytics, and production evidence.").italic = True

    document.add_heading("Live Verification Snapshot", level=1)
    for item in [
        "Database: tijara_dev",
        "Localization: Pakistan country, PKR currency, GST 18% sales tax verified",
        "Demo users: seeded from docs/TEST_CREDENTIALS.csv",
        "Screenshot source: docs/screenshots/INDEX.md",
    ]:
        add_bullet(document, item)

    document.add_heading("Demo Roles", level=1)
    add_credentials_table(document)

    document.add_page_break()
    document.add_heading("Screenshot Walkthrough", level=1)
    for index, figure in enumerate(FIGURES, 1):
        document.add_heading(f"{index}. {figure['title']}", level=2)
        document.add_paragraph(figure["caption"])
        add_picture(document, figure["path"])
        add_caption(document, f"Figure {index}: {figure['title']} captured from live local Odoo.")
        for step in figure["steps"]:
            add_number(document, step)
        if index in {2, 4, 7}:
            document.add_page_break()

    document.add_section(WD_SECTION_START.NEW_PAGE)
    document.add_heading("Workflow Checklist", level=1)
    for title, steps in WORKFLOW_SECTIONS:
        document.add_heading(title, level=2)
        for step in steps:
            add_bullet(document, step)

    document.add_heading("Production Readiness Notes", level=1)
    for item in [
        "Physical hardware certification still needs real printer, scanner, drawer, scale, and display evidence.",
        "FBR live operation still needs certified-provider credentials and compliance sign-off.",
        "Payment providers still need PSP certification, settlement reconciliation, refunds, and chargeback sign-off.",
        "Full staging browser E2E should run against seeded users and real staging URLs before customer deployment.",
        "Monitoring, alerting, restore drills, load tests, and security scans should be attached to release sign-off.",
    ]:
        add_bullet(document, item)

    document.save(DOCX_OUT)


if __name__ == "__main__":
    write_markdown()
    write_docx()
    print(MARKDOWN_OUT)
    print(DOCX_OUT)
